from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib import messages
from django.db import transaction
from django.db.models import Q, Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from apps.auditoria.services import registrar
from apps.cuentas import permisos
from apps.cuentas.models import Usuario

from . import reglas
from .forms import DiaFormSet, FiltroReporteForm, PedidoForm
from .models import Habilitacion, RegistroHoras

DIAS = ["Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom"]
CERO = Decimal(0)


def _fecha(texto, defecto):
    try:
        return date.fromisoformat(texto) if texto else defecto
    except ValueError:
        return defecto


def _url_dia(fecha):
    return f"{reverse('horas:mio')}?dia={fecha.isoformat()}"


def resumen_semana(usuario, lunes):
    """Filas (proyecto · tipo · normal/extra) × días de la semana, con totales por día."""
    dias = [lunes + timedelta(days=i) for i in range(7)]
    feriados = reglas.feriados_en(dias[0], dias[-1])
    registros = RegistroHoras.objects.filter(usuario=usuario, fecha__range=(dias[0], dias[-1])).select_related("obra")
    filas = {}
    totales = {d: {"normal": CERO, "extra": CERO} for d in dias}
    for r in registros:
        clave = (r.obra_id, r.tipo, r.extra)
        fila = filas.setdefault(clave, {"obra": r.obra, "tipo": r.get_tipo_display(), "extra": r.extra,
                                        "horas": {d: CERO for d in dias}, "total": CERO})
        fila["horas"][r.fecha] += r.horas
        fila["total"] += r.horas
        totales[r.fecha]["extra" if r.extra else "normal"] += r.horas
    filas = sorted(filas.values(), key=lambda f: (f["obra"] is None, f["obra"].nombre if f["obra"] else "", f["tipo"], f["extra"]))
    for f in filas:
        f["celdas"] = [f["horas"][d] for d in dias]
    cols = [{"fecha": d, "nombre": DIAS[i], "normal": totales[d]["normal"], "extra": totales[d]["extra"],
             "total": totales[d]["normal"] + totales[d]["extra"],
             "contratadas": reglas.horas_contratadas(usuario, d, feriados),
             "feriado": feriados.get(d), "pagadas": reglas.jornada_feriado(usuario, d, feriados)}
            for i, d in enumerate(dias)]
    return {"filas": filas, "dias": cols,
            "total_normal": sum((c["normal"] for c in cols), CERO), "total_extra": sum((c["extra"] for c in cols), CERO),
            "total_feriado": sum((c["pagadas"] for c in cols), CERO)}


@permisos.solo_personal
def mio(request):
    """Mi Time Report: la semana y la carga del día elegido (hoy, o un día pasado habilitado)."""
    u = request.user
    hoy = reglas.hoy()
    dia = min(_fecha(request.GET.get("dia"), hoy), hoy)
    editable = reglas.puede_editar(u, dia)
    formset = None
    if editable:
        formset = DiaFormSet(request.POST or None, usuario=u, fecha=dia)
        if request.method == "POST":
            if formset.is_valid():
                with transaction.atomic():
                    guardados = formset.save()
                total = sum((r.horas for r in guardados), CERO)
                if dia != hoy:
                    registrar("horas.carga_tardia", f"Cargó {total:g} h del {dia:%d/%m/%Y} con habilitación", objeto=u)
                messages.success(request, f"Horas del {dia:%d/%m} guardadas: {total:g} h.")
                return redirect(_url_dia(dia))
            messages.error(request, "Revisá la carga: hay datos con errores.")
    elif request.method == "POST":
        messages.error(request, "Ese día ya no se puede editar. Pedí una habilitación.")
        return redirect(_url_dia(dia))

    lunes = reglas.lunes_de(dia)
    pedido = Habilitacion.objects.filter(usuario=u, fecha=dia).first() if dia < hoy else None
    return render(request, "horas/mio.html", {
        "dia": dia, "hoy": hoy, "editable": editable, "formset": formset,
        "semana": resumen_semana(u, lunes), "lunes": lunes,
        "semana_anterior": lunes - timedelta(days=7), "semana_siguiente": lunes + timedelta(days=7),
        "contratadas": reglas.horas_contratadas(u, dia),
        "feriado": reglas.feriado(dia),
        "vigente": reglas.habilitacion_vigente(u, dia) if dia < hoy else None,
        "pedido": pedido, "form_pedido": PedidoForm(),
        "puede_habilitar": permisos.puede_habilitar_horas(u),
        "pendientes": Habilitacion.objects.filter(estado=Habilitacion.Estado.PENDIENTE).count()
        if permisos.puede_habilitar_horas(u) else 0,
    })


@permisos.solo_personal
@require_POST
def repetir(request):
    """Copia las cargas del último día con horas (antes del elegido) al día elegido, si está vacío."""
    u = request.user
    dia = min(_fecha(request.POST.get("dia"), reglas.hoy()), reglas.hoy())
    if not reglas.puede_editar(u, dia):
        messages.error(request, "Ese día no se puede editar.")
    elif RegistroHoras.objects.filter(usuario=u, fecha=dia).exists():
        messages.error(request, "Este día ya tiene horas cargadas.")
    else:
        anterior = RegistroHoras.objects.filter(usuario=u, fecha__lt=dia).order_by("-fecha").values_list("fecha", flat=True).first()
        if not anterior:
            messages.error(request, "No hay días anteriores con horas para repetir.")
        else:
            copias = [RegistroHoras(usuario=u, fecha=dia, obra_id=r.obra_id, tipo=r.tipo, horas=r.horas,
                                    extra=r.extra, nota=r.nota)
                      for r in RegistroHoras.objects.filter(usuario=u, fecha=anterior)]
            RegistroHoras.objects.bulk_create(copias)
            messages.success(request, f"Se copiaron {len(copias)} cargas del {anterior:%d/%m}. Ajustá lo que cambie.")
    return redirect(_url_dia(dia))


@permisos.solo_personal
@require_POST
def pedir(request):
    u = request.user
    dia = _fecha(request.POST.get("dia"), None)
    form = PedidoForm(request.POST)
    if not dia or dia >= reglas.hoy():
        messages.error(request, "Solo se pide habilitación para días pasados.")
    elif not form.is_valid():
        messages.error(request, "Contá brevemente por qué necesitás cargar ese día.")
    elif Habilitacion.objects.filter(usuario=u, fecha=dia, estado=Habilitacion.Estado.PENDIENTE).exists():
        messages.info(request, "Ya hay un pedido pendiente para ese día.")
    else:
        pedido = form.save(commit=False)
        pedido.usuario, pedido.fecha = u, dia
        pedido.save()
        registrar("horas.pedido", f"Pidió habilitación para cargar el {dia:%d/%m/%Y}: {pedido.motivo}", objeto=u)
        messages.success(request, "Pedido enviado. Cuando lo aprueben vas a poder cargar ese día hasta las 23:59.")
    return redirect(_url_dia(dia or reglas.hoy()))


@permisos.solo_personal
def habilitaciones(request):
    permisos.exigir(permisos.puede_habilitar_horas(request.user))
    base = Habilitacion.objects.select_related("usuario", "resuelta_por")
    return render(request, "horas/habilitaciones.html", {
        "pendientes": base.filter(estado=Habilitacion.Estado.PENDIENTE).order_by("pedida"),
        "resueltas": base.exclude(estado=Habilitacion.Estado.PENDIENTE).order_by("-resuelta")[:30],
    })


@permisos.solo_personal
@require_POST
def resolver(request, pk):
    permisos.exigir(permisos.puede_habilitar_horas(request.user))
    pedido = get_object_or_404(Habilitacion, pk=pk, estado=Habilitacion.Estado.PENDIENTE)
    aprobar = request.POST.get("accion") == "aprobar"
    pedido.estado = Habilitacion.Estado.APROBADA if aprobar else Habilitacion.Estado.RECHAZADA
    pedido.resuelta_por, pedido.resuelta = request.user, timezone.now()
    pedido.valida_hasta = reglas.fin_del_dia() if aprobar else None
    pedido.save()
    registrar("horas.habilitacion", f"{'Habilitó' if aprobar else 'Rechazó'} la carga del {pedido.fecha:%d/%m/%Y} "
              f"a {pedido.usuario.nombre_completo}", objeto=pedido.usuario)
    messages.success(request, f"{pedido.usuario.nombre_completo}: "
                              f"{'puede cargar el ' + format(pedido.fecha, '%d/%m') + ' hasta hoy a las 23:59' if aprobar else 'pedido rechazado'}.")
    return redirect("horas:habilitaciones")


# --- Reportes ------------------------------------------------------------------

def _agrupado(registros, agrupar):
    campos = {"obra": ("obra__nombre", "obra__codigo"), "persona": ("usuario__first_name", "usuario__last_name"),
              "tipo": ("tipo",)}[agrupar]
    filas = (registros.values(*campos)
             .annotate(normal=Sum("horas", filter=Q(extra=False)), extra=Sum("horas", filter=Q(extra=True)),
                       total=Sum("horas"))
             .order_by("-total"))
    tipos = dict(RegistroHoras.Tipo.choices)
    resultado = []
    for f in filas:
        if agrupar == "obra":
            nombre = f"{f['obra__nombre']} · {f['obra__codigo']}" if f["obra__nombre"] else "Estudio / sin proyecto"
        elif agrupar == "persona":
            nombre = f"{f['usuario__first_name']} {f['usuario__last_name']}".strip()
        else:
            nombre = tipos.get(f["tipo"], f["tipo"])
        resultado.append({"nombre": nombre, "normal": f["normal"] or CERO, "extra": f["extra"] or CERO,
                          "total": f["total"] or CERO})
    return resultado


def _feriados_por_persona(desde, hasta):
    """Horas pagas por feriados (lunes a viernes) en el período, por persona activa."""
    dias = [f for f in reglas.feriados_en(desde, hasta) if f.weekday() < 5]
    if not dias:
        return {}
    return {f"{u.first_name} {u.last_name}".strip(): Decimal(u.horas_diarias or 0) * len(dias)
            for u in Usuario.objects.filter(tipo=Usuario.Tipo.PERSONAL, is_active=True) if u.horas_diarias}


def _faltantes(desde, hasta):
    """Por persona: días hábiles (hasta ayer) en que cargó menos que su jornada."""
    hasta = min(hasta, reglas.hoy() - timedelta(days=1))
    if hasta < desde:
        return []
    habiles = [desde + timedelta(days=i) for i in range((hasta - desde).days + 1)]
    feriados = reglas.feriados_en(desde, hasta)
    habiles = [d for d in habiles if d.weekday() < 5 and d not in feriados]
    cargado = {(r["usuario_id"], r["fecha"]): r["total"] for r in RegistroHoras.objects.filter(
        fecha__range=(desde, hasta), extra=False).values("usuario_id", "fecha").annotate(total=Sum("horas"))}
    resultado = []
    for u in Usuario.objects.filter(tipo=Usuario.Tipo.PERSONAL, is_active=True).order_by("first_name", "last_name"):
        jornada = Decimal(u.horas_diarias or 0)
        if not jornada:
            continue
        dias = [d for d in habiles if cargado.get((u.pk, d), CERO) < jornada]
        if dias:
            resultado.append({"persona": u, "dias": dias})
    return resultado


@permisos.solo_personal
def reportes(request):
    permisos.exigir(permisos.puede_habilitar_horas(request.user))
    hoy = reglas.hoy()
    form = FiltroReporteForm(request.GET or None, initial={
        "desde": hoy.replace(day=1), "hasta": hoy, "agrupar": "obra"})
    if form.is_bound and form.is_valid():
        desde, hasta, agrupar = form.cleaned_data["desde"], form.cleaned_data["hasta"], form.cleaned_data["agrupar"]
    else:
        desde, hasta, agrupar = hoy.replace(day=1), hoy, "obra"
        if not form.is_bound:
            form = FiltroReporteForm(initial={"desde": desde, "hasta": hasta, "agrupar": agrupar})
    registros = RegistroHoras.objects.filter(fecha__range=(desde, hasta))
    filas = _agrupado(registros, agrupar)
    por_persona = _feriados_por_persona(desde, hasta)
    if agrupar == "persona":
        # La jornada paga de los feriados se suma a cada persona (no se asigna a ningún proyecto).
        indice = {f["nombre"]: f for f in filas}
        for nombre, horas in por_persona.items():
            fila = indice.get(nombre) or {"nombre": nombre, "normal": CERO, "extra": CERO, "total": CERO}
            if nombre not in indice:
                filas.append(fila)
            fila["feriado"] = horas
            fila["total"] += horas
        filas.sort(key=lambda f: -f["total"])
    for f in filas:
        f.setdefault("feriado", CERO)
    totales = {k: sum((f[k] for f in filas), CERO) for k in ("normal", "feriado", "extra", "total")}
    feriados_periodo = [f for f in reglas.feriados_en(desde, hasta).values() if f.fecha.weekday() < 5]
    if request.GET.get("formato") == "xlsx":
        return _excel(registros, filas, totales, desde, hasta, dict(FiltroReporteForm.AGRUPAR)[agrupar])
    return render(request, "horas/reportes.html", {
        "form": form, "filas": filas, "totales": totales, "desde": desde, "hasta": hasta, "agrupar": agrupar,
        "faltantes": _faltantes(desde, hasta), "feriados_periodo": feriados_periodo,
        "horas_feriado_total": sum(por_persona.values(), CERO),
    })


def _excel(registros, filas, totales, desde, hasta, titulo):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Resumen"
    hoja.append([f"Time Report · {titulo} · {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}"])
    hoja["A1"].font = Font(bold=True)
    hoja.append([])
    hoja.append(["", "Normales", "Feriados", "Extras", "Total"])
    for f in filas:
        hoja.append([f["nombre"], float(f["normal"]), float(f["feriado"]), float(f["extra"]), float(f["total"])])
    hoja.append(["Total", float(totales["normal"]), float(totales["feriado"]), float(totales["extra"]),
                 float(totales["total"])])
    hoja.column_dimensions["A"].width = 45

    detalle = libro.create_sheet("Detalle")
    detalle.append(["Fecha", "Persona", "Proyecto u obra", "Código", "Tipo de tarea", "Horas", "Extra", "Nota"])
    for r in registros.select_related("usuario", "obra").order_by("fecha", "usuario__first_name"):
        detalle.append([r.fecha, r.usuario.nombre_completo, r.obra.nombre if r.obra else "Estudio / sin proyecto",
                        r.obra.codigo if r.obra else "", r.get_tipo_display(), float(r.horas),
                        "Sí" if r.extra else "No", r.nota])
    for col, ancho in zip("ABCDEFGH", (12, 24, 36, 14, 18, 8, 7, 40)):
        detalle.column_dimensions[col].width = ancho
    for fila in detalle.iter_rows(min_row=2, max_col=1):
        fila[0].number_format = "DD/MM/YYYY"

    salida = BytesIO()
    libro.save(salida)
    respuesta = HttpResponse(salida.getvalue(),
                             content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    respuesta["Content-Disposition"] = f'attachment; filename="time-report_{desde:%Y%m%d}_{hasta:%Y%m%d}.xlsx"'
    return respuesta


def horas_de_obra(obra):
    """Para la pestaña Horas de un proyecto u obra: totales por persona y por tipo de tarea."""
    registros = RegistroHoras.objects.filter(obra=obra)
    return {
        "horas_por_persona": _agrupado(registros, "persona"),
        "horas_por_tipo": _agrupado(registros, "tipo"),
        "horas_total": registros.aggregate(t=Sum("horas"))["t"] or CERO,
        "horas_extra": registros.filter(extra=True).aggregate(t=Sum("horas"))["t"] or CERO,
    }


# --- Calendario de feriados (Gerencia) -----------------------------------------

@permisos.solo_gerencia
def feriados(request):
    from .forms import FeriadoForm
    from .models import Feriado

    anio = int(request.GET.get("anio") or reglas.hoy().year)
    form = FeriadoForm(request.POST or None)
    if request.method == "POST":
        if form.is_valid():
            f = form.save()
            registrar("horas.feriado_alta", f"Nuevo feriado: {f}")
            messages.success(request, f"Feriado agregado: {f}.")
            return redirect(f"{reverse('horas:feriados')}?anio={f.fecha.year}")
        messages.error(request, "Revisá los datos del feriado.")
    return render(request, "horas/feriados.html", {
        "anio": anio, "form": form,
        "feriados": Feriado.objects.filter(fecha__year=anio),
        "anios": sorted({d.year for d in Feriado.objects.dates("fecha", "year")} | {reglas.hoy().year, reglas.hoy().year + 1}),
    })


@permisos.solo_gerencia
@require_POST
def feriado_cambiar(request, pk):
    from .models import Feriado

    f = get_object_or_404(Feriado, pk=pk)
    if request.POST.get("accion") == "borrar":
        registrar("horas.feriado_baja", f"Se quitó el feriado: {f}")
        f.delete()
        messages.success(request, f"Se quitó «{f.nombre}».")
    else:
        f.aplica = not f.aplica
        f.save(update_fields=["aplica"])
        registrar("horas.feriado_cambio", f"Feriado {'activado' if f.aplica else 'desactivado'}: {f}")
    return redirect(f"{reverse('horas:feriados')}?anio={f.fecha.year}")
