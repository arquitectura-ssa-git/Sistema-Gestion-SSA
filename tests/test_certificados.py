"""Importar certificados de avance desde Excel. Se usa una planilla de ejemplo con el mismo formato que la del
contratista (los certificados reales no se guardan en el repositorio)."""

from datetime import date
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile

from apps.obras import certificados
from apps.obras.models import AvanceTarea, Certificado, TareaObra

from .conftest import asignar, cliente_de

ENCABEZADO = [None, "Nº", " ITEM", " UN.", " CANTIDAD", " PU MO", " PF MO", "Precio final\nMO", "Inc %",
              "ANTERIOR", "ACUM", "ACTUAL", "ANTERIOR", "ACUM", "ACTUAL"]


def fila(codigo, nombre, unidad=None, cantidad=None, pu=None, pf=None, rubro=None, acum=None, acum_monto=None,
         actual=None):
    return [1 if rubro is not None else None, codigo, nombre, unidad, cantidad, pu, pf, rubro, None, None, acum,
            None, None, acum_monto, actual]


def planilla(numero=12, fecha="300926", avance_bases=1, avance_medianera=0.5):
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append([None, "PRESUPUESTO OBRA CASA DE PRUEBA - CONTRATISTA", None, None, None, None, None, None, None,
                 f"CERTIFICADO N° {numero}- {fecha}"])
    hoja.append([None, None, "OBRA: CASA DE PRUEBA- ", "SUP DE PROYECTO", None, 100, "m²"])
    hoja.append([None] * 9 + ["AVANCE [%]", None, None, "AVANCE [$]"])
    hoja.append(ENCABEZADO)
    filas = [
        fila(1, "TAREAS PRELIMINARES", rubro=300),
        fila("1.1", "Obrador", "gl", 1, 100, 100, acum=1, acum_monto=100),
        fila("1.9", "MEDIANERAS"),
        fila("1.9.1", "Picado revoque", "m2", 10, 10, 100, acum=avance_medianera, acum_monto=50),
        fila("1.9.2", "Planchado", "m2", 10, 10, 100, acum=0, acum_monto=0),
        fila(2, "HORMIGON", rubro=700),
        fila("2.1", "Bases", "m³", 2, 300, 600, acum=avance_bases, acum_monto=600 * avance_bases),
        fila("3", "CUBIERTA", rubro=0),
        fila("3.1", "Chapa", "m²", None, None, None, acum=0),
        fila("2.2", "Vigas", "m³", 1, 100, 100, acum=0, acum_monto=0),  # fuera de lugar: va a HORMIGON
        [None, None, "SUBTOTAL", None, None, None, None, 1000, 1, None, None, None, None, 750],
        [None, 1, f"SUBTOTAL  1 - CERTIFICADO N° {numero}- {fecha}"] + [None] * 11 + [80],
        [None, 3, f"FONDO DE REPARO 5 % - CERTIFICADO N° {numero}- {fecha}"] + [None] * 11 + [-4],
        [None, None, f"TOTAL SALDO - DESCONTADO LOS ANTICIPOS - CERTIFICADO N° {numero}- {fecha}"] + [None] * 11 + [76],
    ]
    for f in filas:
        hoja.append(f)
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def archivo(**kwargs):
    return SimpleUploadedFile("certificado.xlsx", planilla(**kwargs))


def test_lee_rubros_items_unidades_y_montos():
    datos = certificados.leer(BytesIO(planilla()))
    assert (datos["numero"], datos["fecha"]) == (12, date(2026, 9, 30))
    assert [r["nombre"] for r in datos["rubros"]] == ["TAREAS PRELIMINARES", "HORMIGON", "CUBIERTA"]
    preliminares, hormigon, cubierta = datos["rubros"]
    assert [h["nombre"] for h in preliminares["hijos"]] == ["Obrador", "Medianeras: Picado revoque",
                                                            "Medianeras: Planchado"]
    assert [h["codigo"] for h in hormigon["hijos"]] == ["2.1", "2.2"]  # el 2.2 vuelve a su rubro
    bases = hormigon["hijos"][0]
    assert (bases["unidad"], bases["cantidad"], bases["monto"], bases["avance"]) == ("m³", Decimal("2.000"), Decimal("600.00"), 100)
    assert (datos["presupuesto"], datos["acumulado"], datos["periodo"], datos["saldo"]) == (
        Decimal("1000.00"), Decimal("750.00"), Decimal("80.00"), Decimal("76.00"))
    assert [c for c, _ in datos["resumen"]] == ["SUBTOTAL 1", "FONDO DE REPARO 5 %", "TOTAL SALDO - DESCONTADO LOS ANTICIPOS"]


def test_las_incidencias_salen_de_los_montos_y_suman_100():
    rubros = certificados.leer(BytesIO(planilla()))["rubros"]
    assert sum(r["incidencia"] for r in rubros) == Decimal(100)
    assert rubros[0]["incidencia"] == Decimal("30.000") and rubros[2]["incidencia"] == Decimal("0.001")  # sin precio
    for r in rubros:
        assert sum(h["incidencia"] for h in r["hijos"]) == Decimal(100)


def test_un_archivo_que_no_es_certificado_se_rechaza():
    libro = openpyxl.Workbook()
    libro.active.append(["cualquier", "cosa"])
    salida = BytesIO()
    libro.save(salida)
    with pytest.raises(ValidationError):
        certificados.leer(BytesIO(salida.getvalue()))


def test_importar_arma_la_obra_y_guarda_el_certificado(dir_, obra):
    respuesta = cliente_de(dir_).post(f"/obras/{obra.pk}/certificados/importar/", {"archivo": archivo()})
    assert respuesta.status_code == 302
    obra.refresh_from_db()
    assert obra.tareas.filter(tipo=TareaObra.Tipo.ITEM).count() == 6
    cert = Certificado.objects.get(obra=obra)
    assert (cert.numero, cert.fecha, cert.saldo) == (12, date(2026, 9, 30), Decimal("76.00"))
    # Carga inicial: lo que ya tenía avance queda como "antes del 30/09", sin inventar cuándo empezó.
    inicial = AvanceTarea.objects.get(tarea__nombre="Bases")
    assert (inicial.anterior, inicial.avance, inicial.fecha) == (None, 100, date(2026, 9, 30))
    assert obra.avance > 0


def test_el_siguiente_certificado_actualiza_el_avance_y_conserva_el_historial(dir_, obra):
    c = cliente_de(dir_)
    c.post(f"/obras/{obra.pk}/certificados/importar/", {"archivo": archivo(avance_medianera=0.5)})
    picado = obra.tareas.get(codigo="1.9.1")
    c.post(f"/obras/{obra.pk}/certificados/importar/",
           {"archivo": archivo(numero=13, fecha="151026", avance_medianera=1)})
    assert obra.tareas.get(codigo="1.9.1").pk == picado.pk  # misma fila: no se perdió nada
    registros = list(AvanceTarea.objects.filter(tarea=picado).values_list("fecha", "anterior", "avance"))
    assert registros == [(date(2026, 9, 30), None, 50), (date(2026, 10, 15), 50, 100)]
    assert obra.certificados.count() == 2


def test_no_se_carga_dos_veces_el_mismo_certificado(dir_, obra):
    c = cliente_de(dir_)
    c.post(f"/obras/{obra.pk}/certificados/importar/", {"archivo": archivo()})
    respuesta = c.post(f"/obras/{obra.pk}/certificados/importar/", {"archivo": archivo()}, follow=True)
    assert "ya está cargado" in respuesta.content.decode()
    assert obra.certificados.count() == 1


def test_importar_y_descargar_requieren_permiso(comun, otro, administracion, obra):
    asignar(comun, obra, responsable=True)
    asignar(otro, obra)
    assert cliente_de(otro).post(f"/obras/{obra.pk}/certificados/importar/", {"archivo": archivo()}).status_code == 403
    assert cliente_de(administracion).post(f"/obras/{obra.pk}/certificados/importar/",
                                           {"archivo": archivo()}).status_code == 302
    cert = obra.certificados.get()
    url = f"/obras/{obra.pk}/certificados/{cert.pk}/archivo/"
    assert cliente_de(administracion).get(url).status_code == 200
    assert cliente_de(comun).get(url).status_code == 403  # responsable de la obra, pero sin permiso de montos


def test_los_proyectos_no_importan_certificados(dir_, proyectos):
    proyecto = proyectos[0]
    c = cliente_de(dir_)
    assert "Importar certificado" not in c.get(f"/obras/{proyecto.pk}/?tab=tareas").content.decode()
    assert c.post(f"/obras/{proyecto.pk}/certificados/importar/", {"archivo": archivo()}).status_code == 403


def planilla_material_y_mo():
    """Otro formato de contratista: columnas de material y de mano de obra; el cierre dice
    «SUBTOTAL - CERTIFICADO N° 04- 091026» y el total, «TOTAL - CERTIFICADO…»."""
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append([None, "PRESUPUESTO OBRA CASA DE PRUEBA - 1° ETAPA"])
    hoja.append([None, None, "OBRA: CASA DE PRUEBA - "])
    hoja.append([None] * 17 + ["AVANCE [%]", None, None, "AVANCE [$]"])
    hoja.append([None, "Nº", " ITEM", " UN.", " CANTIDAD", " PU Material", " PF Material", "Precio final\nMaterial",
                 " PU MO", " PF MO", "Precio final\nMO", "Precio final\nMAT - MO", "Incidencia", "Var\nMAT",
                 "materiales", "MO", "Var\nMO", "ANTERIOR", "ACUM", "ACTUAL", "ANTERIOR", "ACUM", "ACTUAL"])

    def f(codigo, nombre, unidad=None, cant=None, pf_mo=None, rubro_mo=None, rubro_total=None, acum=None,
          acum_monto=None, actual=None):
        return [None, codigo, nombre, unidad, cant, None, 999, None, None, pf_mo, rubro_mo, rubro_total, None,
                None, None, None, None, None, acum, None, None, acum_monto, actual]

    for fila_ in [
        f(1, "TAREAS PRELIMINARES", rubro_mo=300, rubro_total=900),
        f("1.1", "Obrador", "gl", 1, 300, acum=1, acum_monto=300, actual=100),
        f("31", "HORMIGON COMPLEMENTARIO"),
        f("31.1", "Llenado de pilotes", "m³", 10, 700, acum=0.5, acum_monto=350, actual=150),
        [None, None, "SUBTOTAL   - CERTIFICADO N° 04- 091026"] + [None] * 19 + [250],
        [None, None, "FONDO DE REPARO 5 % - CERTIFICADO N° 04- 091026"] + [None] * 19 + [-12.5],
        [None, None, "DEVOLUCION ANTICIPO N°1 -  CUOTA 04"] + [None] * 19 + [-30],
        [None, None, "TOTAL   -  CERTIFICADO N° 04- 091026"] + [None] * 19 + [207.5],
    ]:
        hoja.append(fila_)
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def test_lee_el_formato_con_material_y_mano_de_obra():
    datos = certificados.leer(BytesIO(planilla_material_y_mo()))
    assert (datos["numero"], datos["fecha"]) == (4, date(2026, 10, 9))
    preliminares, hormigon = datos["rubros"]
    assert preliminares["monto"] == Decimal("300.00")  # mano de obra, no material + mano de obra
    assert [(h["codigo"], h["avance"]) for h in hormigon["hijos"]] == [("31.1", 50)]
    assert (datos["presupuesto"], datos["acumulado"], datos["periodo"], datos["saldo"]) == (
        Decimal("1000.00"), Decimal("650.00"), Decimal("250.00"), Decimal("207.50"))
    assert [c for c, _ in datos["resumen"]] == ["SUBTOTAL", "FONDO DE REPARO 5 %", "DEVOLUCION ANTICIPO N°1 - CUOTA 04",
                                                 "TOTAL"]


def test_con_varias_pestanas_usa_la_del_certificado():
    libro = openpyxl.load_workbook(BytesIO(planilla()))
    libro.active.title = "CERTIF N°12"
    presupuesto = libro.create_sheet("PPTO A CLIENTE", 0)  # otra pestaña adelante, sin columna de avance
    presupuesto.append([None, "Nº", " ITEM", " UN.", " CANTIDAD", " PU MO", " PF MO"])
    presupuesto.append([None, "1", "OTRA COSA"])
    salida = BytesIO()
    libro.save(salida)
    datos = certificados.leer(BytesIO(salida.getvalue()))
    assert datos["numero"] == 12
    assert [r["nombre"] for r in datos["rubros"]] == ["TAREAS PRELIMINARES", "HORMIGON", "CUBIERTA"]
